"""Gera data/revisao_datas.xlsx: datas da planilha que parecem ter dia e mes trocados (ou que faltam).

Nao grava nada no banco. Uso: python scripts/revisar_datas.py
"""
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from migracao import datas, leitura, limpeza  # noqa: E402

ARQ_PLANILHA = RAIZ / "data" / "orientacao_adm_interna.xlsx"
SAIDA = RAIZ / "data" / "revisao_datas.xlsx"
HOJE = date.today()
ORDEM = {datas.BAIXA: 0, datas.MEDIA: 1, datas.ALTA: 2}
CORES = {datas.BAIXA: "F8CBAD", datas.MEDIA: "FFE699", datas.ALTA: "C6E0B4"}


def br(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else "—"


def montar():
    a = leitura.ler_acompanhamento(ARQ_PLANILHA)
    a["numero"] = a.index + 2  # linha no Excel
    a = a[~(a.doc.map(limpeza.vazio) & a.desvio.map(limpeza.vazio))]  # tira as linhas pre-preenchidas vazias
    a = a[~a.desvio.map(limpeza.vazio)].copy()

    linhas, corrigidas = datas.linhas_da_planilha((r.numero, r.data) for r in a.itertuples())
    por_linha = {x.linha: x for x in linhas}
    sugestoes = datas.resolver(linhas, HOJE)
    # estabilidade: se a resposta muda conforme quantas vizinhas sao consideradas, a confianca cai para baixa
    for n_viz in (2, 5):
        for numero, s in datas.resolver(linhas, HOJE, n_viz).items():
            base = sugestoes[numero]
            if s.sugerida != base.sugerida and base.confianca != datas.BAIXA:
                sugestoes[numero] = datas.Sugestao(
                    numero, base.sugerida, datas.BAIXA, "a resposta muda conforme as vizinhas consideradas: " + base.motivo)

    def contexto(numero, lado):
        vizinhas = [n for n in range(numero - 2, numero) if n in por_linha] if lado < 0 else \
                   [n for n in range(numero + 1, numero + 3) if n in por_linha]
        return " | ".join(f"L{n}: {br(por_linha[n].data)}" for n in vizinhas) or "—"

    saida = []
    for r in a.itertuples():
        x = por_linha[r.numero]
        sug = None
        if x.data is None:
            sug = datas.sugerir_sem_data(linhas, r.numero)
            tipo = "sem data"
        elif r.numero in corrigidas:
            sug = datas.Sugestao(r.numero, x.data, datas.ALTA, f"erro de digitação corrigido (estava '{r.data}')")
            tipo = "texto com erro de digitação"
        elif x.ambigua:
            sug = sugestoes[r.numero]
            tipo = "data real do Excel (dia até 12)"
        elif x.data > HOJE:
            sug = datas.Sugestao(r.numero, x.data, datas.BAIXA, "data no futuro e não é ambígua: conferir")
            tipo = "texto"
        else:
            continue
        muda = sug.sugerida != x.data
        saida.append({
            "Linha no Excel": r.numero,
            "CNPJ/CPF": r.doc,
            "Prestador (na planilha)": r.nome,
            "Desvio": r.desvio,
            "Tipo da célula": tipo,
            "Data hoje no sistema": br(x.data) if x.data else "(não migrada: sem data)",
            "Data sugerida": br(sug.sugerida),
            "Muda?": "SIM" if muda else "não",
            "Confiança": sug.confianca,
            "Motivo": sug.motivo,
            "2 linhas antes": contexto(r.numero, -1),
            "2 linhas depois": contexto(r.numero, +1),
            "Decisão (preencher se discordar)": "",
        })
    df = pd.DataFrame(saida)
    df["_o"] = df["Confiança"].map(ORDEM)
    return df.sort_values(["_o", "Linha no Excel"]).drop(columns="_o").reset_index(drop=True)


def gravar(df):
    resumo = pd.DataFrame({
        "Item": ["Linhas listadas", "  que mudam de data (Muda? = SIM)", "  que ficam como estão",
                 "Confiança ALTA", "Confiança MÉDIA", "Confiança BAIXA (decisão sua)",
                 "Datas no futuro hoje no sistema"],
        "Quantidade": [len(df), int((df["Muda?"] == "SIM").sum()), int((df["Muda?"] == "não").sum()),
                       int((df["Confiança"] == datas.ALTA).sum()), int((df["Confiança"] == datas.MEDIA).sum()),
                       int((df["Confiança"] == datas.BAIXA).sum()),
                       int(df["Data hoje no sistema"].map(
                           lambda s: s[6:] + s[3:5] + s[:2] > HOJE.strftime("%Y%m%d") if s[:1].isdigit() else False).sum())],
    })
    legenda = pd.DataFrame({"Como ler": [
        "Cada linha é uma data da planilha que pode estar com dia e mês trocados (o Excel leu 12/06 como 6 de dezembro).",
        "'Data hoje no sistema' é o que está gravado agora; 'Data sugerida' é a leitura que combina com as linhas vizinhas.",
        "Filtre 'Muda?' = SIM para ver só o que seria alterado. As datas estão como TEXTO dd/mm/aaaa de propósito.",
        "Comece pelas de confiança BAIXA (vermelhas): é onde a planilha sozinha não decide.",
        "Se discordar da sugestão, escreva a data certa (dd/mm/aaaa) em 'Decisão'. Em branco = aceita a sugestão.",
    ]})
    with pd.ExcelWriter(SAIDA, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Revisão", index=False)
        resumo.to_excel(w, sheet_name="Resumo", index=False)
        legenda.to_excel(w, sheet_name="Resumo", index=False, startrow=len(resumo) + 3)
        ws = w.sheets["Revisão"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        larguras = [9, 18, 34, 34, 26, 20, 14, 8, 11, 56, 38, 38, 26]
        for i, larg in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = larg
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for linha in ws.iter_rows(min_row=2):
            linha[8].fill = PatternFill("solid", fgColor=CORES[linha[8].value])
            for c in linha:
                c.number_format = "@"
                c.alignment = Alignment(vertical="top", wrap_text=c.column in (10, 11, 12))
        r = w.sheets["Resumo"]
        r.column_dimensions["A"].width = 60
        r.column_dimensions["B"].width = 14
        for c in r[1]:
            c.font = Font(bold=True)


if __name__ == "__main__":
    df = montar()
    gravar(df)
    print(f"{len(df)} linhas | muda: {(df['Muda?'] == 'SIM').sum()}")
    print(df["Confiança"].value_counts().to_string())
    print("gerado:", SAIDA)
